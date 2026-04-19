"""
ADMI Financial Model → Dashboard Data Extractor

Reads the Excel financial model and produces a JSON file the HTML dashboard uses.
Extracts BOTH In-Person and Hybrid scenarios by toggling Summary!C5 and recalculating.

Usage:  python extract_data.py <path_to_excel> <output_json>
Default paths: model.xlsx -> data.json
"""

import sys, os, json, shutil, subprocess, warnings
import openpyxl
warnings.filterwarnings('ignore')

# --- CONFIG: exact cell map based on model structure ---------------------

# Diplomas tab: 6 programs across columns (name col, value col)
DIPLOMA_COLS = [
    ('Digital Content Creation Diploma',          'B', 'D'),
    ('Music Production Diploma',                  'G', 'I'),
    ('Sound Engineering Diploma',                 'L', 'N'),
    ('Graphic Design Diploma',                    'Q', 'S'),
    ('Film and Television Production Diploma',    'V', 'X'),
    ('Animation & Motion Graphics Diploma',      'AA', 'AC'),
]

# Certificates tab: 8 programs
CERT_COLS = [
    ('AI Adoption & Digital Transformation Certificate',  'B',  'D'),
    ('Multimedia Certificate',                            'G',  'I'),
    ('Photography Certificate',                           'L',  'N'),
    ('Music Production and Sound Engineering Certificate','Q',  'S'),
    ('Video Production Certificate',                      'V',  'X'),
    ('Graphic Design Certificate',                       'AA',  'AC'),
    ('Digital Marketing Certificate',                    'AF',  'AH'),
    ('Data Analytics and Visualization Certificate',     'AK',  'AM'),
]

# Row layout (same for Diplomas and Certificates tabs)
REVENUE_ROWS = {
    'Course Price':            9,
    'Internship Fee':         10,
    'Application Fee':        11,
    'Insurance Fee':          12,
    'Graduation Fee':         13,
    'Equipment Fee':          14,
    'Accreditation onboarding': 15,
}
ENROLLMENT_ROW = 16
TOTAL_REVENUE_ROW = 17  # Non-Adjusted Revenue

COGS_ROWS = {
    'Direct staff costs':         26,
    'Rent and repairs':           27,
    'Equipment costs':            28,
    'Compute and SW':             29,
    "Students' accreditation":    30,
    'Other - academics':          31,
    'Marketing and advertisements':32,
}
TOTAL_COGS_ROW = 34

# Summary tab OpEx cells
OPEX_CELLS = {
    'Administrative Personnel Costs': 'C28',
    'Insurance Cost':                 'C29',
    'Compliance Cost':                'C30',
    'Professional Fees':              'C31',
    'Directors fees':                 'C32',
    'Office Support expenses':        'C33',
    'Rent and repairs (fixed)':       'C34',
    'Compute and SW (fixed)':         'C35',
    'Equipment costs (fixed)':        'C36',
    'Direct staff costs (fixed)':     'C37',
    'Marketing and advertisement':    'C38',
    'Academics - other (fixed)':      'C39',
    'Provision for Bad Debt':         'C40',
}

# ---------------------------------------------------------------------------

def num(v):
    """Safely coerce cell value to a number. Errors / None / 'TBD' -> 0."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s == '' or s.startswith('#') or s.upper() == 'TBD':
        return 0.0
    try:
        return float(s.replace(',', ''))
    except ValueError:
        return 0.0


def extract_scenario(xlsx_path, scenario_label):
    """Open the data-only workbook and pull all values for whichever scenario
    is currently set in Summary!C5. Returns a dict of programs + opex."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    def read_program_block(ws, name, name_col, val_col):
        revenue = {k: num(ws[f'{val_col}{r}'].value) for k, r in REVENUE_ROWS.items()}
        cogs = {k: num(ws[f'{val_col}{r}'].value) for k, r in COGS_ROWS.items()}
        enrollment = num(ws[f'{val_col}{ENROLLMENT_ROW}'].value)
        total_rev = num(ws[f'{val_col}{TOTAL_REVENUE_ROW}'].value)
        total_cogs = num(ws[f'{val_col}{TOTAL_COGS_ROW}'].value)
        return {
            'name': name,
            'revenue_lines': revenue,
            'cogs_lines': cogs,
            'enrollment': enrollment,
            'total_revenue_sheet': total_rev,  # what the sheet says (sanity)
            'total_cogs_sheet':    total_cogs,
        }

    programs = []
    ws_d = wb['Diplomas']
    for name, name_col, val_col in DIPLOMA_COLS:
        p = read_program_block(ws_d, name, name_col, val_col)
        p['type'] = 'Diploma'
        programs.append(p)

    ws_c = wb['Certificates']
    for name, name_col, val_col in CERT_COLS:
        p = read_program_block(ws_c, name, name_col, val_col)
        p['type'] = 'Certificate'
        programs.append(p)

    ws_s = wb['Summary']
    opex = {k: num(ws_s[cell].value) for k, cell in OPEX_CELLS.items()}

    return {
        'scenario': scenario_label,
        'programs': programs,
        'opex': opex,
    }


def set_scenario_and_recalc(xlsx_path, work_path, scenario):
    """Copy file, set Summary!C5 to the scenario, recalc via Excel (Mac)."""
    shutil.copy(xlsx_path, work_path)
    abs_path = os.path.abspath(work_path)

    script = f'''
tell application "Microsoft Excel"
    set wb to open workbook workbook file name POSIX file "{abs_path}"
    set ws to sheet "Summary" of wb
    set value of cell "C5" of ws to "{scenario}"
    calculate
    delay 5
    calculate
    delay 2
    save wb
    close wb saving yes
end tell
'''
    subprocess.run(['osascript', '-e', script], capture_output=True, check=False)


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else 'model.xlsx'
    out = sys.argv[2] if len(sys.argv) > 2 else 'data.json'

    if not os.path.exists(src):
        print(f"ERROR: {src} not found"); sys.exit(1)

    print(f"Reading: {src}")
    scenarios = {}
    for label in ['In-Person', 'Hybrid']:
        work = f'/tmp/_admi_{label.replace("-","").lower()}.xlsx'
        print(f"  Extracting {label} scenario...")
        set_scenario_and_recalc(src, work, label)
        scenarios[label] = extract_scenario(work, label)

    with open(out, 'w') as f:
        json.dump(scenarios, f, indent=2)
    print(f"Wrote: {out}")

    # Quick sanity print
    for label, s in scenarios.items():
        rev = sum(p['total_revenue_sheet'] for p in s['programs'])
        cogs = sum(p['total_cogs_sheet'] for p in s['programs'])
        opex = sum(s['opex'].values())
        print(f"  {label}: Revenue={rev:,.0f}  COGS={cogs:,.0f}  OpEx={opex:,.0f}")


if __name__ == '__main__':
    main()
